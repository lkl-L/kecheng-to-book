# -*- coding: utf-8 -*-
"""把全书 markdown 重新排成"书籍版式"docx。

版式要点
  · A4，四周大留白（上 2.6 / 下 2.4 / 左右 3.0 cm）
  · 正文 宋体 11.5pt，行距 1.6 倍，段后 5pt，首行缩进 2 字符，两端对齐
  · 原文块 楷体 + 左右缩进；【原文】【讲解】标记作小标题
  · 卷标题、讲标题均另起一页（page_break_before）
  · 封面页 / 编者说明页 / 目录页 各自独立成页
  · 页眉（书名 + 细线）、页脚（居中页码），首页不显示
"""
import os
import re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BASE = os.path.dirname(os.path.abspath(__file__))
MD = os.path.join(BASE, "【书名】.md")
OUT = os.path.join(BASE, "【书名】.docx")
OUT_FALLBACK = os.path.join(BASE, "【书名】（精排版）.docx")

SONG, HEI, KAI = "宋体", "黑体", "楷体"
EN = "Times New Roman"
BODY = 11.5

BOOK_TITLE = "【书名】"


# ---------- 底层工具 ----------

def style_font(style, cn, en=EN, size=None, bold=None, color=(0, 0, 0)):
    style.font.name = en
    rPr = style.element.get_or_add_rPr()
    rf = rPr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rPr.insert(0, rf)
    rf.set(qn("w:ascii"), en)
    rf.set(qn("w:hAnsi"), en)
    rf.set(qn("w:eastAsia"), cn)
    if size is not None:
        style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    if color is not None:
        style.font.color.rgb = RGBColor(*color)


def run_font(run, cn, en=EN, size=BODY, bold=False):
    run.font.name = en
    run.font.size = Pt(size)
    run.font.bold = bold
    rPr = run._element.get_or_add_rPr()
    rf = rPr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rPr.insert(0, rf)
    rf.set(qn("w:ascii"), en)
    rf.set(qn("w:hAnsi"), en)
    rf.set(qn("w:eastAsia"), cn)


def indent_chars(par, n):
    """首行缩进 n 个字符（用 w:firstLineChars，随字号自适应）"""
    ind = par._p.get_or_add_pPr().get_or_add_ind()
    ind.set(qn("w:firstLineChars"), str(int(n * 100)))


def para_border_bottom(par, sz=6, color="999999"):
    pPr = par._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    b.set(qn("w:val"), "single")
    b.set(qn("w:sz"), str(sz))
    b.set(qn("w:space"), "2")
    b.set(qn("w:color"), color)
    pBdr.append(b)
    pPr.append(pBdr)


def add_field(par, instr, size=10):
    """插入域代码（如 PAGE / NUMPAGES）"""
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), instr)
    r = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    rf = OxmlElement("w:rFonts")
    rf.set(qn("w:ascii"), EN)
    rf.set(qn("w:hAnsi"), EN)
    rf.set(qn("w:eastAsia"), SONG)
    rPr.append(rf)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(size * 2)))
    rPr.append(sz)
    r.append(rPr)
    t = OxmlElement("w:t")
    t.text = "1"
    r.append(t)
    fld.append(r)
    par._p.append(fld)


# ---------- 文档骨架 ----------

def setup_page(doc):
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21.0), Cm(29.7)      # A4
    s.top_margin, s.bottom_margin = Cm(2.6), Cm(2.4)
    s.left_margin, s.right_margin = Cm(3.0), Cm(3.0)
    s.header_distance, s.footer_distance = Cm(1.5), Cm(1.4)
    s.different_first_page_header_footer = True

    # 页眉：书名 + 细线
    hp = s.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = hp.add_run(BOOK_TITLE)
    run_font(r, SONG, size=9)
    hp.paragraph_format.space_after = Pt(2)
    para_border_bottom(hp, sz=4, color="BFBFBF")

    # 页脚：居中页码
    fp = s.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run("— ")
    run_font(r, SONG, size=10)
    add_field(fp, "PAGE  \\* MERGEFORMAT", size=10)
    r = fp.add_run(" —")
    run_font(r, SONG, size=10)


def setup_styles(doc):
    st = doc.styles["Normal"]
    style_font(st, SONG, size=BODY)
    pf = st.paragraph_format
    pf.line_spacing = 1.6
    pf.space_after = Pt(5)
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    for name, size in (("Heading 1", 24), ("Heading 2", 22), ("Heading 3", 15)):
        s = doc.styles[name]
        style_font(s, HEI, size=size, bold=True)
        s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        s.paragraph_format.line_spacing = 1.3


def blank(doc, n=1, size=BODY):
    for _ in range(n):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.2
        run_font(p.add_run(""), SONG, size=size)


def add_body(doc, text, cite=False):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing = 1.6
    pf.space_after = Pt(5)
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if cite:
        pf.left_indent = Cm(0.8)
        pf.right_indent = Cm(0.8)
        pf.line_spacing = 1.5
    else:
        indent_chars(p, 2)
    run_font(p.add_run(text), KAI if cite else SONG, size=BODY)
    return p


def add_marker(doc, text):
    """【原文】【讲解】标记行 → 小标题式"""
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(12)
    pf.space_after = Pt(6)
    pf.line_spacing = 1.3
    pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run_font(p.add_run(text), HEI, size=BODY + 0.5, bold=True)
    return p


# ---------- 内容解析 ----------

def parse(md_text):
    """返回 (h1, head_short[], intro[], items[])；items 元素为 (kind, text)"""
    h1 = None
    head_short, intro, items = [], [], []
    started_volume = False
    pending_head_paras = []

    for raw in md_text.split("\n"):
        ln = raw.strip()
        if not ln:
            continue
        if ln.startswith("# ") and not ln.startswith("## "):
            h1 = ln[2:].strip()
            continue
        if ln.startswith("## ") and not ln.startswith("### "):
            started_volume = True
            items.append(("vol", ln[3:].strip()))
            continue
        if ln.startswith("### "):
            started_volume = True
            items.append(("lec", ln[4:].strip()))
            continue
        if not started_volume:
            pending_head_paras.append(ln)
            continue
        if re.fullmatch(r"【[^】]{1,6}】", ln):
            items.append(("mark", ln))
        else:
            items.append(("p", ln))

    # 封面短句（<=30 字）与编者说明分流
    for t in pending_head_paras:
        (head_short if len(t) <= 30 else intro).append(t)
    return h1, head_short, intro, items


# ---------- 主流程 ----------

def main():
    md_text = open(MD, encoding="utf-8").read()
    h1, head_short, intro, items = parse(md_text)

    doc = Document()
    setup_page(doc)
    setup_styles(doc)

    # ===== 封面页 =====
    blank(doc, 7)
    p = doc.add_paragraph(style="Heading 1")
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(28)
    p.paragraph_format.line_spacing = 1.25
    run_font(p.add_run("【书名】"), HEI, size=28, bold=True)
    for t in head_short:
        pp = doc.add_paragraph()
        pp.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pp.paragraph_format.space_after = Pt(8)
        run_font(pp.add_run(t), KAI, size=12)
    blank(doc, 6)
    pp = doc.add_paragraph()
    pp.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_font(pp.add_run("讲师　讲授　　整理者　整理"), SONG, size=11)
    doc.add_page_break()

    # ===== 编者说明 =====
    p = doc.add_paragraph()
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(20)
    run_font(p.add_run("编 者 说 明"), HEI, size=17, bold=True)
    for t in intro:
        add_body(doc, t)
    doc.add_page_break()

    # ===== 目录 =====
    p = doc.add_paragraph()
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(20)
    run_font(p.add_run("目　　录"), HEI, size=17, bold=True)

    # 按卷分组：[卷名, 讲1, 讲2, ...]
    groups = []
    for kind, text in items:
        if kind == "vol":
            groups.append([text])
        elif kind == "lec":
            if groups:
                groups[-1].append(text)
    for g in groups:
        for i, t in enumerate(g):
            pp = doc.add_paragraph()
            if i == 0:                       # 卷名
                pp.paragraph_format.space_before = Pt(10)
                pp.paragraph_format.space_after = Pt(3)
                run_font(pp.add_run(t), HEI, size=12.5, bold=True)
            else:                            # 讲名
                pp.paragraph_format.space_before = Pt(1)
                pp.paragraph_format.space_after = Pt(1)
                pp.paragraph_format.left_indent = Cm(0.9)
                run_font(pp.add_run(t), SONG, size=11)
            pp.paragraph_format.line_spacing = 1.4
            pp.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            # 一卷的目录块整块不拆页：卷名与其后各讲绑定（组内最后一条不绑，免与下一卷粘连）
            pp.paragraph_format.keep_with_next = (i < len(g) - 1)

    # ===== 正文 =====
    in_cite = False
    first_vol = True
    for kind, text in items:
        if kind == "vol":
            # 卷封面页：另起一页 + 上部留白 + 居中卷名 + 装饰线
            p = doc.add_paragraph()
            p.paragraph_format.page_break_before = True
            p.paragraph_format.space_after = Pt(0)
            run_font(p.add_run(""), SONG, size=BODY)
            blank(doc, 5)
            pv = doc.add_paragraph(style="Heading 2")
            pv.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pv.paragraph_format.space_after = Pt(14)
            pv.paragraph_format.line_spacing = 1.4
            run_font(pv.add_run(text), HEI, size=23, bold=True)
            pd = doc.add_paragraph()
            pd.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run_font(pd.add_run("— — —"), SONG, size=12)
            in_cite = False
            first_vol = False
        elif kind == "lec":
            p = doc.add_paragraph(style="Heading 3")
            p.paragraph_format.page_break_before = True
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(22)
            p.paragraph_format.line_spacing = 1.35
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run_font(p.add_run(text), HEI, size=16, bold=True)
            para_border_bottom(p, sz=6, color="808080")
            in_cite = False
        elif kind == "mark":
            add_marker(doc, text)
            in_cite = (text == "【原文】")   # 正文块标记，按你的稿子改
        else:
            add_body(doc, text, cite=in_cite)

    # 保存（原文件被编辑器占用时降级为新文件名）
    try:
        doc.save(OUT)
        print("OK ->", OUT)
    except PermissionError:
        doc.save(OUT_FALLBACK)
        print("OK（原文件被占用）->", OUT_FALLBACK)

    # 统计
    print("段落数:", len(doc.paragraphs))
    print("卷:", sum(1 for k, _ in items if k == "vol"),
          " 讲:", sum(1 for k, _ in items if k == "lec"),
          " 正文段:", sum(1 for k, _ in items if k == "p"))


if __name__ == "__main__":
    main()
