# -*- coding: utf-8 -*-
"""把各书 md 排成统一的书籍版式 docx（v2 版式，风水书系列）。

v2 相对 v1 的改进
  1. 四级标题体系统一：编/卷(H1) → 章/节(H2) → 小节「一、」(H3) → 目「（一）」(H4)
     v1 里「一、」「（一）」是纯文本，被当成正文（首行缩进 2 字符、宋体），和正文分不开。
  2. 封面书名不再占用 Heading 1 样式 → 不再混进目录（v1 目录第一条是书名）。
  3. 目录只收编/章两级（\\o "1-2"），小节不进目录，目录干净。
  4. 正文改为中文书籍惯例：首行缩进 2 字符 + 行距 1.5 + 段后 3pt（v1 是 1.6 + 5pt，三重叠加显松散）。
  5. 标题带 keep_with_next，标题不会孤零零留在页尾；标题不缩进。
  6. 页边距对称（上下 2.54 / 左右 2.8），页面利用率更高。
  7. 新增「图注」样式（（图x-y　……））：楷体小一号 + 左右缩进，与正文明显区分。

标题层级自适应（同 v1）
  · 一层：# 书名 + ## 节                    → 节另起一页
  · 两层：# 书名 + # 编 + ## 节             → 编扉页另起一页、节另起一页
  · 三层：# 书名 + ## 卷 + ### 讲           → 卷扉页另起一页、讲另起一页
用法：python format_books.py [书名1 书名2 ...]   （不带参数=全部）
"""
import os
import re
import sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BASE = os.path.dirname(os.path.abspath(__file__))
SONG, HEI, KAI = "宋体", "黑体", "楷体"
EN = "Times New Roman"

BODY = 11.5          # 正文字号
LINE = 1.50          # 正文行距
PARA_AFTER = 3       # 正文段后间距(pt)

ALL_BOOKS = [
    "经典课实录",
    "赋文课实录",
    "图文课实录",
    "日课课实录",
    "实战课实录",
    "微课堂实录",
]

# 各书的讲授者（书名 → 讲授者）；留空则封面只署整理者
LECTURER = {b: "讲师" for b in ALL_BOOKS}

CENTER = WD_ALIGN_PARAGRAPH.CENTER
LEFT = WD_ALIGN_PARAGRAPH.LEFT
JUST = WD_ALIGN_PARAGRAPH.JUSTIFY

# 标题样式规格： (cn, size, align, space_before, space_after, line_spacing)
HEAD_SPECS = [
    ("Heading 1", HEI, 20.0, CENTER, 0, 20, 1.30),   # 编 / 卷
    ("Heading 2", HEI, 15.5, CENTER, 0, 14, 1.30),   # 章 / 节
    ("Heading 3", HEI, 13.0, LEFT, 16, 8, 1.25),     # 小节「一、」
    ("Heading 4", HEI, 11.5, LEFT, 12, 6, 1.25),     # 目「（一）」
]

# 小节标题识别：形如「一、xxx」「（一）xxx」，且短、不以句末标点收尾
RE_SUB_CN = re.compile(r"^[一二三四五六七八九十]{1,3}、\S")
RE_SUB_PAR = re.compile(r"^（[一二三四五六七八九十]{1,3}）\S")
SUB_MAXLEN = 34
END_PUNCT = ("。", "！", "？", "；", "…")

# 图注：（图1-1　……）
RE_CAPTION = re.compile(r"^（图\s?\d")


# ---------------- 底层工具 ----------------

def style_font(style, cn, en=EN, size=None, bold=None, color=(0, 0, 0), italic=False):
    style.font.name = en
    style.font.italic = italic          # 一律正体：Word 内置样式默认可能带斜体，必须显式关掉
    rPr = style.element.get_or_add_rPr()
    rf = rPr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rPr.insert(0, rf)
    rf.set(qn("w:ascii"), en)
    rf.set(qn("w:hAnsi"), en)
    rf.set(qn("w:eastAsia"), cn)
    if size is not None:
        style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    if color is not None:
        style.font.color.rgb = RGBColor(*color)


def run_font(run, cn, en=EN, size=BODY, bold=False, italic=False):
    run.font.name = en
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic            # 一律正体
    rPr = run._element.get_or_add_rPr()
    rf = rPr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rPr.insert(0, rf)
    rf.set(qn("w:ascii"), en)
    rf.set(qn("w:hAnsi"), en)
    rf.set(qn("w:eastAsia"), cn)


def indent_chars(par, n):
    ind = par._p.get_or_add_pPr().get_or_add_ind()
    ind.set(qn("w:firstLineChars"), str(int(n * 100)))


def para_border_bottom(par, sz=6, color="999999"):
    pPr = par._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    b.set(qn("w:val"), "single")
    b.set(qn("w:sz"), str(sz))
    b.set(qn("w:space"), "4")
    b.set(qn("w:color"), color)
    pBdr.append(b)
    pPr.append(pBdr)


def add_field(par, instr, size=10, text="1", dirty=False):
    """插入域（PAGE / TOC 等）。dirty=True 让 Word 打开时提示/自动更新该域。"""
    if dirty:
        par._p.append(_dirty_flag())
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), instr)
    r = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    rf = OxmlElement("w:rFonts")
    rf.set(qn("w:ascii"), EN)
    rf.set(qn("w:hAnsi"), EN)
    rf.set(qn("w:eastAsia"), SONG)
    rPr.append(rf)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(size * 2)))
    rPr.append(sz)
    r.append(rPr)
    t = OxmlElement("w:t")
    t.text = text
    r.append(t)
    fld.append(r)
    par._p.append(fld)


def _dirty_flag():
    r = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    f = OxmlElement("w:fldChar")
    f.set(qn("w:fldCharType"), "begin")
    f.set(qn("w:dirty"), "true")
    r.append(rPr)
    r.append(f)
    return r


def settings_update_fields(doc):
    """在 settings.xml 里打开 updateFields，Word 打开文档时会自动更新所有域。"""
    st = doc.settings.element
    for tag in ("w:updateFields",):
        el = st.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            st.append(el)
        el.set(qn("w:val"), "true")


# ---------------- 文档骨架 ----------------

def setup_page(doc, book_title):
    s = doc.sections[0]
    s.page_width, s.page_height = Cm(21.0), Cm(29.7)
    s.top_margin, s.bottom_margin = Cm(2.54), Cm(2.54)
    s.left_margin, s.right_margin = Cm(2.8), Cm(2.8)
    s.header_distance, s.footer_distance = Cm(1.5), Cm(1.4)
    s.different_first_page_header_footer = True

    hp = s.header.paragraphs[0]
    hp.alignment = CENTER
    run_font(hp.add_run(book_title), SONG, size=9)
    hp.paragraph_format.space_after = Pt(2)
    para_border_bottom(hp, sz=4, color="BFBFBF")

    fp = s.footer.paragraphs[0]
    fp.alignment = CENTER
    run_font(fp.add_run("— "), SONG, size=10)
    add_field(fp, "PAGE  \\* MERGEFORMAT", size=10)
    run_font(fp.add_run(" —"), SONG, size=10)


def setup_styles(doc):
    st = doc.styles["Normal"]
    style_font(st, SONG, size=BODY)
    pf = st.paragraph_format
    pf.line_spacing = LINE
    pf.space_after = Pt(PARA_AFTER)
    pf.alignment = JUST
    # 中文断行：允许标点溢出、禁止行首标点
    for tag in ("w:kinsoku", "w:overflowPunct", "w:autoSpaceDE", "w:autoSpaceDN"):
        el = OxmlElement(tag)
        el.set(qn("w:val"), "true")
        st.element.get_or_add_pPr().append(el)

    for name, cn, size, align, before, after, ls in HEAD_SPECS:
        s = doc.styles[name]
        style_font(s, cn, size=size, bold=True, color=(0, 0, 0))
        p = s.paragraph_format
        p.alignment = align
        p.space_before = Pt(before)
        p.space_after = Pt(after)
        p.line_spacing = ls
        p.keep_with_next = True

    # 全书一律正体：把 Word 内置样式（Heading 4/6/7/9、Subtitle、Quote、Emphasis…）
    # 自带的斜体全部显式关掉，避免「（一）xxx」这一级目显示成斜体。
    for s in doc.styles:
        try:
            s.font.italic = False
        except (AttributeError, ValueError, KeyError):
            pass


def blank(doc, n=1, size=BODY):
    for _ in range(n):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.2
        run_font(p.add_run(""), SONG, size=size)


def add_body(doc, text, cite=False):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing = LINE
    pf.space_after = Pt(PARA_AFTER)
    pf.alignment = JUST
    if cite:
        pf.left_indent = Cm(1.0)
        pf.right_indent = Cm(1.0)
        pf.line_spacing = 1.40
        pf.space_after = Pt(4)
    else:
        indent_chars(p, 2)
    run_font(p.add_run(text), KAI if cite else SONG, size=BODY)
    return p


def add_caption(doc, text):
    """图注：（图x-y　……）楷体小一号、左右缩进、无首行缩进。"""
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing = 1.35
    pf.space_before = Pt(2)
    pf.space_after = Pt(6)
    pf.left_indent = Cm(0.7)
    pf.right_indent = Cm(0.7)
    pf.alignment = JUST
    run_font(p.add_run(text), KAI, size=BODY - 0.5)
    return p


def add_sub(doc, text, level):
    """小节标题（H3=「一、」 H4=「（一）」）；不缩进、不首行缩进。"""
    p = doc.add_paragraph(style="Heading %d" % level)
    p.paragraph_format.alignment = LEFT
    run_font(p.add_run(text), HEI, size=(13.0 if level == 3 else 11.5), bold=True)
    return p


def add_marker(doc, text):
    """【原文】【讲解】一类的内容标记：居中黑体，与正文区分。"""
    p = doc.add_paragraph(style="Heading 4")
    pf = p.paragraph_format
    pf.alignment = CENTER
    pf.space_before = Pt(12)
    pf.space_after = Pt(6)
    pf.keep_with_next = True
    run_font(p.add_run(text), HEI, size=BODY + 0.5, bold=True)
    return p


# ---------------- 解析 ----------------

def is_sub(ln):
    """判断是否为「一、」「（一）」形式的小节标题。"""
    if len(ln) > SUB_MAXLEN:
        return 0
    if ln.endswith(END_PUNCT):
        return 0
    if RE_SUB_CN.match(ln):
        return 3
    if RE_SUB_PAR.match(ln):
        return 4
    return 0


def parse(md_text):
    """返回 book, cover[], intro[], seq[]；seq 元素 = [kind, text, notes]"""
    lines = md_text.split("\n")
    has_h3 = any(l.startswith("### ") for l in lines)
    book = None
    cover, intro = [], []
    seq = []
    state = "cover"
    cur_vol = None
    stats = {"sub3": 0, "sub4": 0, "mark": 0, "cite": 0, "cap": 0}

    for raw in lines:
        ln = raw.strip()
        if not ln:
            continue
        if ln.startswith("### "):
            seq.append(["sec", ln[4:].strip(), []])
            state, cur_vol = "body", None
        elif ln.startswith("## "):
            if has_h3:                                  # 二层是卷
                cur_vol = ["vol", ln[3:].strip(), []]
                seq.append(cur_vol)
                state = "front"
            else:                                       # 二层就是节
                seq.append(["sec", ln[3:].strip(), []])
                state, cur_vol = "body", None
        elif ln.startswith("# "):
            if book is None:
                book = ln[2:].strip()
                state = "cover"
            else:                                       # 后续 H1 = 编
                cur_vol = ["vol", ln[2:].strip(), []]
                seq.append(cur_vol)
                state = "front"
        else:
            if ln.startswith("> "):                     # 引文块（reparagraph.py 提行）
                seq.append(["cite", ln[2:].strip(), []])
                stats["cite"] += 1
            elif state == "cover":
                (cover if len(ln) <= 30 else intro).append(ln)
            elif state == "front":
                cur_vol[2].append(ln)
            elif re.fullmatch(r"【[^】]{1,8}】", ln):
                seq.append(["mark", ln, []])
                stats["mark"] += 1
            elif RE_CAPTION.match(ln):
                seq.append(["cap", ln, []])
                stats["cap"] += 1
            else:
                lv = is_sub(ln)
                if lv:
                    seq.append(["sub%d" % lv, ln, []])
                    stats["sub%d" % lv] += 1
                else:
                    seq.append(["p", ln, []])
    return book, cover, intro, seq, stats


# ---------------- 生成 ----------------

def book_plain(book):
    """去掉书名号，用于查表与页眉。"""
    return re.sub(r"[《》]", "", book or "").strip()


def build(name):
    md_path = os.path.join(BASE, name + ".md")
    out_path = os.path.join(BASE, name + ".docx")
    if not os.path.exists(md_path):
        print("!! 缺 md:", md_path)
        return None

    md_text = open(md_path, encoding="utf-8").read()
    book, cover, intro, seq, stats = parse(md_text)

    doc = Document()
    setup_page(doc, book)
    setup_styles(doc)
    settings_update_fields(doc)

    # ===== 封面（书名用普通段落，不进目录）=====
    blank(doc, 7)
    p = doc.add_paragraph()
    p.paragraph_format.alignment = CENTER
    p.paragraph_format.space_after = Pt(28)
    p.paragraph_format.line_spacing = 1.25
    n = len(book)
    size = 30 if n <= 10 else (26 if n <= 14 else 22)
    run_font(p.add_run(book), HEI, size=size, bold=True)
    for t in cover:
        p2 = doc.add_paragraph()
        p2.paragraph_format.alignment = CENTER
        p2.paragraph_format.space_after = Pt(8)
        run_font(p2.add_run(t), KAI, size=12)
    blank(doc, 6)
    pp = doc.add_paragraph()
    pp.paragraph_format.alignment = CENTER
    teacher = LECTURER.get(book_plain(book), "")
    line = ("%s　讲授　　" % teacher if teacher else "") + "整理者　整理"
    run_font(pp.add_run(line), SONG, size=11)
    doc.add_page_break()

    # ===== 编者说明（有长段才出）=====
    if intro:
        p = doc.add_paragraph()
        p.paragraph_format.alignment = CENTER
        p.paragraph_format.space_after = Pt(20)
        run_font(p.add_run("编 者 说 明"), HEI, size=17, bold=True)
        for t in intro:
            add_body(doc, t)
        doc.add_page_break()

    # ===== 目录（自动目录域：只收编/章两级，小节不进目录）=====
    p = doc.add_paragraph()
    p.paragraph_format.alignment = CENTER
    p.paragraph_format.space_after = Pt(20)
    run_font(p.add_run("目　　录"), HEI, size=17, bold=True)
    ptoc = doc.add_paragraph()
    ptoc.paragraph_format.space_after = Pt(0)
    add_field(ptoc, 'TOC \\o "1-2" \\h \\z \\u', size=11,
              text="（页码更新中……在 Word 中右键「更新域」即可刷新）", dirty=True)

    # ===== 正文 =====
    in_cite = False
    for it in seq:
        kind, text, notes = it
        if kind == "vol":
            p = doc.add_paragraph()
            p.paragraph_format.page_break_before = True
            p.paragraph_format.space_after = Pt(0)
            run_font(p.add_run(""), SONG, size=BODY)
            blank(doc, 5 if notes else 7)
            pv = doc.add_paragraph(style="Heading 1")
            pv.paragraph_format.alignment = CENTER
            pv.paragraph_format.space_before = Pt(0)
            pv.paragraph_format.space_after = Pt(12)
            pv.paragraph_format.line_spacing = 1.35
            run_font(pv.add_run(text), HEI, size=20, bold=True)
            pd = doc.add_paragraph()
            pd.paragraph_format.alignment = CENTER
            run_font(pd.add_run("— — —"), SONG, size=12)
            if notes:                            # 编前说明排在同一张扉页下半部
                blank(doc, 2)
                for t in notes:
                    pn = doc.add_paragraph()
                    pn.paragraph_format.line_spacing = 1.5
                    pn.paragraph_format.space_after = Pt(4)
                    pn.paragraph_format.left_indent = Cm(1.2)
                    pn.paragraph_format.right_indent = Cm(1.2)
                    pn.paragraph_format.alignment = JUST
                    run_font(pn.add_run(t), KAI, size=11)
            in_cite = False
        elif kind == "sec":
            p = doc.add_paragraph(style="Heading 2")
            p.paragraph_format.page_break_before = True
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(14)
            p.paragraph_format.line_spacing = 1.30
            p.paragraph_format.alignment = CENTER
            run_font(p.add_run(text), HEI, size=15.5, bold=True)
            in_cite = False
        elif kind == "mark":
            add_marker(doc, text)
            in_cite = (text == "【经文】")
        elif kind in ("sub3", "sub4"):
            add_sub(doc, text, 3 if kind == "sub3" else 4)
            in_cite = False
        elif kind == "cite":
            add_body(doc, text, cite=True)
            in_cite = False
        elif kind == "cap":
            add_caption(doc, text)
        else:
            add_body(doc, text, cite=in_cite)

    doc.save(out_path)
    n_vol = sum(1 for it in seq if it[0] == "vol")
    n_sec = sum(1 for it in seq if it[0] == "sec")
    n_p = sum(1 for it in seq if it[0] == "p")
    print("OK  %-30s 编%-2d 章/节%-3d 小节%-3d 标记%-3d 引文%-4d 图注%-3d 正文段%-5d 段落总%-5d %s"
          % (name, n_vol, n_sec, stats["sub3"], stats["mark"], stats["cite"],
             stats["cap"], n_p, len(doc.paragraphs),
             "%d KB" % (os.path.getsize(out_path) // 1024)))
    return True


def main():
    targets = sys.argv[1:] or ALL_BOOKS
    for name in targets:
        build(name)


if __name__ == "__main__":
    main()
